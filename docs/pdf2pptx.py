#!/usr/bin/env python3
"""Build a pixel-perfect PowerPoint from the rendered pitch deck.

Each page of PITCH.pdf becomes a full-bleed image on a 16:9 slide, so the .pptx
is visually identical to the branded PDF (metropolis coral+charcoal) and immune
to font substitution on whoever opens it. Not text-editable by design — PITCH.md
stays the source of truth, and PITCH-editable.pptx is the editable variant.

Usage: pdf2pptx.py <png-dir> <out.pptx>   (png-dir holds slide-NN.png pages)
"""
import sys
import glob
import os
from pptx import Presentation
from pptx.util import Inches

png_dir, out = sys.argv[1], sys.argv[2]
pages = sorted(glob.glob(os.path.join(png_dir, "*.png")))
if not pages:
    sys.exit("no PNG pages found in " + png_dir)

prs = Presentation()
prs.slide_width = Inches(13.333)   # 16:9, matches the beamer deck
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]       # the "Blank" layout

for page in pages:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(page, 0, 0,
                             width=prs.slide_width, height=prs.slide_height)

prs.save(out)
print("wrote %s with %d slides" % (out, len(pages)))
